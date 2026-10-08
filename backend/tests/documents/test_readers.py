import base64
import io
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch
from app.documents import DocumentReadError, create_document_reader
from app.documents.formats import MEDIA, CODE_FORMATS, EXTENSIONS, FILENAMES
from app.documents.isolation import read_isolated
from app.rag.parsers import NativeDocumentParser


def image_bytes():
    from PIL import Image
    buffer = io.BytesIO()
    Image.new('RGB', (20, 20), 'white').save(buffer, 'PNG')
    return buffer.getvalue()


def fake_vision(available=True):
    vision = Mock()
    vision.capabilities.return_value = SimpleNamespace(image_ocr=available, pdf_page_ocr=available,
        image_formats=('png', 'jpg') if available else ())
    vision.asset_revisions.return_value = ()
    vision.extract_image_text.return_value = 'Image-only explanation\n'
    return vision


class DocumentReaderTests(unittest.TestCase):
    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.root = Path(folder.name)
        self.vision = fake_vision()
        self.reader = create_document_reader(vision=self.vision, isolated=False)

    def write(self, name, data):
        path = self.root / name
        path.write_bytes(data.encode() if isinstance(data, str) else data)
        return path

    def test_txt_keeps_previous_paragraph_and_indentation_behavior(self):
        result = self.reader.read(self.write('notes.txt', '\ufeffFirst paragraph\n\n    return wrapper\n\n'))
        self.assertEqual(result.blocks, [{'text':'First paragraph'}, {'text':'    return wrapper'}])
        self.assertEqual(result.media_type, 'text/plain')
        self.assertEqual(result.warnings, [])

    def test_coding_files_keep_every_character_and_language(self):
        text = '# comment\r\n\r\ndef café():\r\n    return "hello"\r\n'
        for extension in ('py', 'ts', 'json', 'yaml', 'ps1', 'vue'):
            with self.subTest(extension=extension):
                content = self.reader.read(self.write('source.' + extension, text))
                self.assertEqual(content.text, text)
                self.assertEqual(content.blocks[0]['language'], extension)
                self.assertEqual(content.media_type, MEDIA[extension])

    def test_all_source_extensions_and_filenames_preserve_source(self):
        text = '# source \u03bb\r\n\r\n    keep_every_character();\r\n'
        cases = [('source' + extension.upper(), kind) for extension, kind in EXTENSIONS.items() if kind in CODE_FORMATS]
        cases += [(name, kind) for name, kind in FILENAMES.items()]
        cases += [('Dockerfile.dev', 'dockerfile'), ('.env.production', 'env'), ('Containerfile.test', 'dockerfile')]
        for name, kind in cases:
            with self.subTest(filename=name):
                content = self.reader.read(self.write(name, text))
                self.assertEqual(content.text, text)
                self.assertEqual(content.format, kind)
                self.assertEqual(content.blocks[0]['language'], kind)

    def test_extensionless_source_still_rejects_binary_and_invalid_utf8(self):
        for name, data in [('Dockerfile', b'\x00binary'), ('Makefile', b'\xffinvalid'), ('.env.local', b'\x00binary')]:
            with self.subTest(filename=name), self.assertRaises(DocumentReadError):
                self.reader.read(self.write(name, data))

    def test_html_keeps_head_inline_script_style_pre_and_comments(self):
        js = '\n  const text = "<p>hello & world</p>";\n'
        css = '\n  .card { color: red; }\n'
        source = '<html><head><style>' + css + '</style></head><body><h1>Python</h1><p>Read <b>this</b> &amp; learn.</p><script>' + js + '</script><pre>    return wrapper\n</pre><!-- internal note --><div hidden>hidden source text</div></body></html>'
        result = self.reader.read(self.write('notes.html', source))
        by_language = {block.get('language'): block['text'] for block in result.blocks if block.get('language')}
        self.assertEqual(by_language['js'], js)
        self.assertEqual(by_language['css'], css)
        self.assertEqual(by_language['html'], '<!-- internal note -->')
        self.assertIn('Read this & learn.', result.text)
        self.assertIn('hidden source text', result.text)
        self.assertIn('    return wrapper\n', [block['text'] for block in result.blocks])

    def test_html_embedded_images_use_vision_once_without_fetching_urls(self):
        encoded = base64.b64encode(image_bytes()).decode()
        html = f'<h1>Images</h1><img src="data:image/png;base64,{encoded}"><img src="data:image/png;base64,{encoded}"><img src="https://example.invalid/private.png" alt="Remote diagram"><script src="https://example.invalid/app.js"></script>'
        result = self.reader.read(self.write('images.htm', html))
        self.vision.extract_image_text.assert_called_once()
        self.assertEqual(len([block for block in result.blocks if block.get('method') == 'image_ocr']), 2)
        self.assertIn('Remote diagram', result.text)
        self.assertTrue(any('not fetched' in warning for warning in result.warnings))
        self.assertEqual(result.format, 'html')

    def test_html_nested_template_and_json_script_are_not_discarded(self):
        source = '<div>' * 1500 + '<template>Template content</template><script type="application/ld+json">\r\n {"name": "Mentra"}\r\n</script>' + '</div>' * 1500
        result = self.reader.read(self.write('nested.html', source))
        self.assertIn('Template content', result.text)
        self.assertEqual(next(block['text'] for block in result.blocks if block.get('language') == 'json'), '\r\n {"name": "Mentra"}\r\n')

    def test_docx_body_and_table_images_use_vision_with_heading_provenance(self):
        from docx import Document
        document = Document()
        document.add_heading('Python', level=1)
        document.add_paragraph('Native notes')
        document.add_picture(io.BytesIO(image_bytes()))
        document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run().add_picture(io.BytesIO(image_bytes()))
        path = self.root / 'notes.docx'
        document.save(path)
        result = self.reader.read(path)
        extracted = [block for block in result.blocks if block.get('method') == 'image_ocr']
        self.assertEqual(len(extracted), 2)
        self.assertTrue(all(block['heading_path'] == ['Python'] for block in extracted))
        self.vision.extract_image_text.assert_called_once()
        self.assertIn('Native notes', result.text)

    def test_pptx_grouped_pictures_and_native_text_share_slide_provenance(self):
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = 'Python'
        group = slide.shapes.add_group_shape()
        group.shapes.add_picture(io.BytesIO(image_bytes()), Inches(1), Inches(1))
        path = self.root / 'notes.pptx'
        presentation.save(path)
        result = self.reader.read(path)
        block = next(block for block in result.blocks if block.get('method') == 'image_ocr')
        self.assertEqual(block['slide'], 1)
        self.assertEqual(block['heading_path'], ['Python'])
        self.vision.extract_image_text.assert_called_once()

    def test_unavailable_vision_warns_instead_of_losing_native_office_text(self):
        from docx import Document
        document = Document()
        document.add_paragraph('Keep my native text')
        document.add_picture(io.BytesIO(image_bytes()))
        path = self.root / 'notes.docx'
        document.save(path)
        vision = fake_vision(False)
        result = create_document_reader(vision=vision, isolated=False).read(path)
        self.assertIn('Keep my native text', result.text)
        self.assertTrue(any('OCR is unavailable' in warning for warning in result.warnings))
        vision.extract_image_text.assert_not_called()

    def test_broken_embedded_images_warn_and_keep_native_content(self):
        self.vision.extract_image_text.side_effect = OSError('private image failure')
        data = base64.b64encode(image_bytes()).decode()
        result = self.reader.read(self.write('notes.html', f'<p>Native</p><img src="data:image/png;base64,{data}">'))
        self.assertEqual(result.text, 'Native')
        self.assertTrue(any('could not be read' in warning for warning in result.warnings))
        self.assertFalse(any('private image failure' in warning for warning in result.warnings))

    def test_content_signatures_override_misleading_extensions(self):
        from testing.rag_real import build_pdf
        result = self.reader.read(self.write('misleading.txt', build_pdf('Python decorators wrap functions.')))
        self.assertEqual(result.format, 'pdf')

    def test_invalid_utf8_and_binary_source_files_fail_explicitly(self):
        for data in (b'\xffbad', b'\x00binary'):
            with self.subTest(data=data), self.assertRaises(DocumentReadError):
                self.reader.read(self.write('bad.py', data))

    def test_limits_fail_before_starting_a_reader(self):
        path = self.write('notes.txt', 'source')
        for options in ({'max_pages':0}, {'timeout':float('nan')}, {'max_bytes':1}, {'max_pages':True}):
            with self.subTest(options=options), self.assertRaises(DocumentReadError):
                self.reader.read(path, **options)

    def test_new_format_can_be_composed_without_rewriting_the_facade(self):
        engine = SimpleNamespace(formats=('txt',), available=Mock(return_value=True),
            read=Mock(return_value={'blocks':[{'text':'custom adapter'}], 'warnings':[]}))
        result = create_document_reader(vision=self.vision, engines=[engine], isolated=False).read(self.write('notes.txt', 'source'))
        self.assertEqual(result.text, 'custom adapter')
        engine.read.assert_called_once()
        with self.assertRaises(ValueError):
            create_document_reader(vision=self.vision)

    def test_rag_adapter_only_translates_document_errors(self):
        reader = Mock()
        reader.detect.side_effect = DocumentReadError('safe document error')
        reader.parse.side_effect = DocumentReadError('safe document error')
        parser = NativeDocumentParser(reader)
        from app.rag.errors import ExtractionError
        with self.assertRaisesRegex(ExtractionError, 'safe document error'):
            parser.detect(Path('file'), 'file.txt')
        with self.assertRaisesRegex(ExtractionError, 'safe document error'):
            parser.parse(Path('file'), 'text/plain', 200, 120)

    def test_default_public_api_reads_in_the_shared_subprocess(self):
        path = self.write('script.py', '    return wrapper\n')
        result = create_document_reader().read(path)
        self.assertEqual(result.text, '    return wrapper\n')
        self.assertEqual(result.blocks[0]['language'], 'py')

    def test_standalone_import_does_not_initialize_consuming_modules(self):
        code = 'import sys; from app.documents import create_document_reader; create_document_reader(); assert not any(m.startswith(("app.rag", "app.learner", "app.langchain", "app.db")) for m in sys.modules)'
        subprocess.run([sys.executable, '-c', code], check=True, capture_output=True, timeout=20)

    def test_isolation_rejects_invalid_child_output(self):
        process = Mock(returncode=0)
        process.communicate.return_value = (b'not-json', b'private error')
        with patch('app.documents.isolation.subprocess.Popen', return_value=process), self.assertRaises(DocumentReadError):
            read_isolated(Path('notes'), 'txt', 200, 1)

    def test_isolation_timeout_cancels_child_tree(self):
        process = Mock(pid=1234)
        process.communicate.side_effect = [subprocess.TimeoutExpired('reader', 1), (b'', b'')]
        with patch('app.documents.isolation.subprocess.Popen', return_value=process), \
            patch('app.documents.isolation.os.killpg', create=True) as killpg, patch('app.documents.isolation.subprocess.run') as taskkill, \
            self.assertRaisesRegex(DocumentReadError, 'time limit'):
            read_isolated(Path('notes'), 'txt', 200, 1)
        self.assertEqual(killpg.call_count + taskkill.call_count, 1)
