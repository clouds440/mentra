import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from types import SimpleNamespace
from app.documents import create_document_reader


def parse_native(path, kind, max_pages):
    return create_document_reader(isolated=False).read_in_process(path, kind, max_pages)


class ParserStructureTests(unittest.TestCase):
    def test_docx_nested_headings_and_code_indentation_survive(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'notes.docx'
            document = Document()
            document.add_heading('Python', level=1)
            document.add_heading('Decorators', level=2)
            document.add_paragraph('    return wrapper')
            document.save(path)
            blocks = parse_native(path, 'docx', 200)['blocks']
            self.assertEqual(blocks[-1]['heading_path'], ['Python', 'Decorators'])
            self.assertEqual(blocks[-1]['text'], '    return wrapper')

    def test_pptx_grouped_text_is_not_silently_lost(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'notes.pptx'
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            group = slide.shapes.add_group_shape()
            box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
            box.text = 'Grouped decorator examples'
            presentation.save(path)
            blocks = parse_native(path, 'pptx', 200)['blocks']
            self.assertEqual(blocks[0]['text'], 'Grouped decorator examples')
            self.assertEqual(blocks[0]['slide'], 1)

    def test_short_native_pdf_text_does_not_require_ocr(self):
        from testing.rag_real import build_pdf
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'short.pdf'
            path.write_bytes(build_pdf('Notes'))
            vision = SimpleNamespace(capabilities=lambda: SimpleNamespace(image_formats=(), image_ocr=False, pdf_page_ocr=False), asset_revisions=lambda: ())
            parsed = create_document_reader(vision=vision, isolated=False).read_in_process(path, 'pdf', 200)
            self.assertEqual(parsed['blocks'][0]['text'].strip(), 'Notes')
            self.assertEqual(parsed['blocks'][0]['method'], 'native')
