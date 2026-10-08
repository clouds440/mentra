"""Real standalone document facade and Office/HTML image-to-Vision verification."""
import base64
import io
import json
import tempfile
from pathlib import Path
from app.documents import create_document_reader
from testing.rag_real import formats


def main():
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    reader = create_document_reader()
    fixtures = formats()
    image = fixtures[4][1]
    doc = Document()
    doc.add_heading('Python', level=1)
    doc.add_paragraph('Native office notes')
    doc.add_picture(io.BytesIO(image))
    doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run().add_picture(io.BytesIO(image))
    doc_buffer = io.BytesIO()
    doc.save(doc_buffer)
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = 'Python'
    slide.shapes.add_group_shape().shapes.add_picture(io.BytesIO(image), Inches(1), Inches(1))
    slide_buffer = io.BytesIO()
    slides.save(slide_buffer)
    html = ('<h1>Python</h1><script>const decorator = "wrapper";</script><style>.python { color: red; }</style>'
        '<img src="data:image/png;base64,' + base64.b64encode(image).decode() + '">').encode()
    fixtures += [('docx', doc_buffer.getvalue()), ('pptx', slide_buffer.getvalue()), ('html', html),
        ('py', b'def decorator():\r\n    return "wrapper"\r\n')]
    fixtures += [('ts', b'export const count: number = 1;\r\n'),
        ('tsx', b'export const View = () => <div>Hello</div>;\r\n'),
        ('js', b'// Keep comments\r\nexport const count = 1;\r\n'),
        ('jsx', b'export const View = () => <div>Hello</div>;\r\n')]
    report = dict(module='mentra documents', reader_revision=reader.version, checks=[])
    with tempfile.TemporaryDirectory() as folder:
        for number, (kind, data) in enumerate(fixtures):
            path = Path(folder) / f'fixture-{number}.{kind}'
            path.write_bytes(data)
            result = reader.read(path)
            assert result.text.strip(), kind
            if number >= 8 and kind in ('docx', 'pptx', 'html'):
                image_blocks = [block for block in result.blocks if block.get('method') == 'image_ocr']
                assert image_blocks and all(block['text'].lower() == 'python decorators wrap functions.' for block in image_blocks), result
            if kind == 'html':
                assert 'const decorator = "wrapper";' in result.text and '.python { color: red; }' in result.text
            if kind in ('py', 'ts', 'tsx', 'js', 'jsx'):
                assert result.text == data.decode()
            report['checks'].append(dict(format=kind, fixture=number, blocks=len(result.blocks),
                methods=[block.get('method', 'native') for block in result.blocks],
                languages=[block['language'] for block in result.blocks if block.get('language')], warnings=result.warnings))
    print(json.dumps(report, indent=2))


if __name__ == '__main__':
    main()
