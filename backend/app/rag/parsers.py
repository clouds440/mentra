"""Bounded document adapters using shared Mentra Vision in a killable subprocess."""
import importlib.util
import json
import subprocess
import sys
import os
import signal
import tempfile
import zipfile
import re
from importlib.metadata import version as package_version, PackageNotFoundError
from pathlib import Path
from app.rag.errors import ExtractionError
from app.vision import VisionError, create_vision_service

_vision = create_vision_service()

PARSER_VERSION = 'native-local-ocr-v1'
MEDIA = {'txt': 'text/plain', 'pdf': 'application/pdf', 'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
         'pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation', 'png': 'image/png', 'jpg': 'image/jpeg'}


def capabilities():
    formats = ['txt']
    for kind, module in (('pdf','pypdf'), ('docx','docx'), ('pptx','pptx')):
        if importlib.util.find_spec(module):
            formats.append(kind)
    vision = _vision.capabilities()
    formats += list(vision.image_formats)
    return dict(formats=formats, image_ocr=vision.image_ocr, scanned_pdf_ocr=vision.pdf_page_ocr)


def detect(path, filename):
    with path.open('rb') as source:
        prefix = source.read(16)
    suffix = Path(filename).suffix.lower()
    if prefix.startswith(b'%PDF-'):
        kind = 'pdf'
    elif prefix.startswith(b'\x89PNG\r\n\x1a\n'):
        kind = 'png'
    elif prefix.startswith(b'\xff\xd8\xff'):
        kind = 'jpg'
    elif zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            if len(archive.infolist()) > 10000 or sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ExtractionError('Compressed material exceeds the extraction limit.')
            names = set(archive.namelist())
            kind = 'docx' if 'word/document.xml' in names else 'pptx' if 'ppt/presentation.xml' in names else ''
    elif suffix == '.txt':
        kind = 'txt'
        try:
            path.read_text(encoding='utf-8-sig')
        except UnicodeError as exc:
            raise ExtractionError('Text files must use UTF-8 encoding.') from exc
        if b'\x00' in prefix:
            raise ExtractionError('Binary files cannot be uploaded as text.')
    else:
        kind = ''
    if kind not in capabilities()['formats']:
        raise ExtractionError('This file format is unsupported on this server.')
    return kind


def ocr_image(path):
    try:
        return _vision.extract_image_text(path)
    except VisionError as exc:
        raise ExtractionError(str(exc)) from exc


def ocr_pdf_page(path, page):
    try:
        return _vision.extract_pdf_page_text(path, page)
    except VisionError as exc:
        raise ExtractionError(str(exc)) from exc


def parse_native(path, kind, max_pages):
    blocks, warnings = [], []
    if kind == 'txt':
        text = path.read_text(encoding='utf-8-sig')
        if '\x00' in text:
            raise ExtractionError('Binary content cannot be parsed as text.')
        blocks = [dict(text=p) for p in text.split('\n\n') if p.strip()]
    elif kind == 'docx':
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        heading = []
        for block in Document(path).iter_inner_content():
            if isinstance(block, Paragraph):
                if block.style and block.style.name.startswith('Heading'):
                    match = re.search(r'(\d+)$', block.style.name)
                    level = int(match.group(1)) if match else 1
                    heading = heading[:max(0, level - 1)] + [block.text]
                blocks.append(dict(text=block.text, heading_path=heading[:]))
            elif isinstance(block, Table):
                blocks.append(dict(text='\n'.join(' | '.join(c.text for c in row.cells) for row in block.rows), heading_path=heading[:]))
        warnings.append('Embedded images and handwritten annotations in office documents are not extracted.')
    elif kind == 'pptx':
        from pptx import Presentation
        presentation = Presentation(path)
        if len(presentation.slides) > max_pages:
            raise ExtractionError('Presentation exceeds the slide limit.')
        def shapes_in(shapes):
            for shape in shapes:
                if hasattr(shape, 'shapes'):
                    yield from shapes_in(shape.shapes)
                else:
                    yield shape
        for number, slide in enumerate(presentation.slides, 1):
            heading = [slide.shapes.title.text] if slide.shapes.title else []
            for shape in shapes_in(slide.shapes):
                text = '\n'.join(' | '.join(c.text for c in row.cells) for row in shape.table.rows) if shape.has_table else shape.text if shape.has_text_frame else ''
                if text.strip():
                    blocks.append(dict(text=text, slide=number, heading_path=heading))
        warnings.append('Speaker notes and text inside slide images are not extracted.')
    elif kind == 'pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ExtractionError('Password-protected PDFs are unsupported. Upload an unlocked copy.')
        if len(reader.pages) > max_pages:
            raise ExtractionError('PDF exceeds the page limit.')
        ocr_capabilities = capabilities()
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text(extraction_mode='layout') or ''
            images = page.images
            method = 'native'
            if len(text.strip()) < 25 and (not text.strip() or len(images)):
                if not ocr_capabilities['scanned_pdf_ocr']:
                    if len(images):
                        raise ExtractionError('This PDF needs OCR. The server OCR tools are unavailable.')
                    warnings.append(f'Page {number} has no readable native text and OCR is unavailable; check the original.')
                    blocks.append(dict(text=text, page=number, method=method))
                    continue
                text, method = ocr_pdf_page(path, number), 'ocr'
                warnings.append(f'Page {number} was read with OCR; verify equations and handwriting.')
            if not text.strip():
                warnings.append(f'Page {number} has no readable text; check the original for missing content.')
            blocks.append(dict(text=text, page=number, method=method))
            # Mixed pages retain native text and OCR only their embedded images.
            # We report page provenance without inventing image bounding boxes.
            if method == 'native' and ocr_capabilities['image_ocr']:
                for image in images:
                    with tempfile.TemporaryDirectory() as folder:
                        image_path = Path(folder) / 'image.png'
                        image.image.save(image_path, format='PNG')
                        extra = ocr_image(image_path).strip()
                    if extra and extra.casefold() not in text.casefold():
                        blocks.append(dict(text=extra, page=number, method='image_ocr'))
                        warnings.append(f'Image text on page {number} was read with OCR; verify it against the source.')
            elif method == 'native' and len(images):
                warnings.append(f'Image text on page {number} was not extracted because OCR is unavailable.')
    elif kind in ('png','jpg'):
        blocks = [dict(text=ocr_image(path), method='ocr')]
        warnings.append('This image was read with English OCR; verify handwriting and equations.')
    if sum(len(block['text']) for block in blocks) > 5_000_000:
        raise ExtractionError('Extracted text exceeds the processing limit. Split this material into smaller files.')
    if any('\x00' in block['text'] for block in blocks):
        warnings.append('Unsupported null characters were removed during extraction.')
        for block in blocks:
            block['text'] = block['text'].replace('\x00', '')
    return dict(blocks=blocks, warnings=list(dict.fromkeys(warnings)))


def parse_isolated(path, kind, max_pages, timeout):
    process = subprocess.Popen([sys.executable, '-m', 'app.rag.parsers', str(path), kind, str(max_pages)],
        stdout=subprocess.PIPE, stderr=subprocess.PIPE, start_new_session=os.name != 'nt')
    try:
        output, _error = process.communicate(timeout=timeout)
    except subprocess.TimeoutExpired as exc:
        if os.name != 'nt':
            try:
                os.killpg(process.pid, signal.SIGKILL)
            except ProcessLookupError:
                pass
        else:
            subprocess.run(['taskkill', '/PID', str(process.pid), '/T', '/F'], capture_output=True, check=False)
            process.kill()
        process.communicate()
        raise ExtractionError('Parsing exceeded the time limit. Split the material into smaller files.') from exc
    if process.returncode:
        # Only explicitly sanitized domain messages cross the subprocess boundary.
        try:
            error = json.loads(output).get('error')
        except (ValueError, UnicodeError):
            error = None
        raise ExtractionError(error or 'The file could not be parsed. Check it and upload an unlocked, valid copy.')
    return json.loads(output)


class NativeDocumentParser:
    def __init__(self):
        assets = []
        for package in ('pypdf', 'python-docx', 'python-pptx'):
            try:
                revision = package_version(package)
            except PackageNotFoundError:
                revision = 'unavailable'
            assets.append(f'{package}={revision}')
        assets.extend(f'{asset}={revision}' for asset, revision in _vision.asset_revisions())
        self.version = PARSER_VERSION + ':' + ';'.join(assets)

    def detect(self, path: Path, filename: str) -> str:
        return MEDIA[detect(path, filename)]

    def capabilities(self) -> dict:
        return dict(**capabilities(), parser_revision=self.version)

    def parse(self, path: Path, media_type: str, max_pages: int, timeout: int) -> dict:
        kind = next((kind for kind, media in MEDIA.items() if media == media_type), None)
        if kind is None:
            raise ExtractionError('This media type is unsupported by the document parser.')
        return parse_isolated(path, kind, max_pages, timeout)


if __name__ == '__main__':
    try:
        # Linux worker child has a hard virtual-memory/CPU limit; parent timeout
        # remains active on every platform. No network-capable parsers are used.
        if sys.platform == 'linux':
            import resource
            resource.setrlimit(resource.RLIMIT_AS, (1024**3, 1024**3))
            resource.setrlimit(resource.RLIMIT_CPU, (120, 120))
        parsed = parse_native(Path(sys.argv[1]), sys.argv[2], int(sys.argv[3]))
        print(json.dumps(parsed, ensure_ascii=True))
    except ExtractionError as exc:
        print(json.dumps({'error': str(exc)}))
        sys.exit(1)
    except Exception:
        print(json.dumps({'error': 'The file could not be parsed. Upload a valid, unlocked copy.'}))
        sys.exit(1)
