import importlib.util
from app.documents.errors import DocumentReadError
from app.documents.readers.images import EmbeddedImages


from app.core.logging import workflow_logger

@workflow_logger.connect_module(default_outcome='success')
class PPTXReader:
    formats = ('pptx',)
    def __init__(self, vision):
        self.vision = vision

    def available(self):
        return importlib.util.find_spec('pptx') is not None

    def read(self, path, kind, max_pages):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        presentation = Presentation(path)
        if len(presentation.slides) > max_pages:
            raise DocumentReadError('Presentation exceeds the slide limit.')
        blocks, warnings = [], []
        embedded = EmbeddedImages(self.vision)
        def shapes_in(shapes):
            for shape in shapes:
                if hasattr(shape, 'shapes'):
                    yield from shapes_in(shape.shapes)
                else:
                    yield shape
        for number, slide in enumerate(presentation.slides, 1):
            heading = [slide.shapes.title.text] if slide.shapes.title else []
            for shape in shapes_in(slide.shapes):
                text = '\n'.join(' | '.join(c.text for c in row.cells) for row in shape.table.rows) if shape.has_table else shape.text if shape.has_text_frame else ''
                if text.strip():
                    blocks.append(dict(text=text, slide=number, heading_path=heading))
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    extra = embedded.read(shape.image.blob, warnings, f'on slide {number}')
                    if extra:
                        blocks.append(dict(text=extra, slide=number, heading_path=heading, method='image_ocr'))
                        warnings.append(f'Image text on slide {number} was read with English OCR; verify it against the source.')
        warnings.append('Speaker notes, charts and drawing text may be incomplete; inspect the original.')
        return dict(blocks=blocks, warnings=warnings)
